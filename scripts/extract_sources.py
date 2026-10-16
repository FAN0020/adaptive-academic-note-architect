#!/usr/bin/env python3
"""Extract located academic text; visual review is still required."""
import argparse
import json
from pathlib import Path
import sys
from zipfile import ZipFile
from docx_utils import NS, W, parse, paragraphs, visible_text, style_id, sha256, write_json

VERSION = '2.0.0'

def extract(path):
    suffix = path.suffix.lower()
    units, warnings = [], []
    if suffix in {'.txt', '.md', '.markdown', '.csv', '.vtt', '.srt'}:
        text = path.read_text(encoding='utf-8-sig')
        units = [{'location': f'line {i}', 'text': s} for i,s in enumerate(text.splitlines(),1)]
    elif suffix == '.pdf':
        from pypdf import PdfReader
        r = PdfReader(path)
        if r.is_encrypted and not r.decrypt(''): raise ValueError('Password-protected PDF')
        for i,p in enumerate(r.pages,1):
            text = p.extract_text() or ''
            units.append({'location': f'PDF page {i}', 'text': text})
            if len(text.strip()) < 30: warnings.append(f'Page {i}: little/no text; inspect visually or OCR')
        warnings.append('PDF text may omit diagrams, equations, images or reading order; visual review required where relevant')
    elif suffix == '.docx':
        with ZipFile(path) as z:
            root = parse(z.read('word/document.xml'))
            for i,p in enumerate(paragraphs(root)):
                units.append({'location': f'word/document.xml paragraph {i}', 'paragraph_index': i,
                              'style': style_id(p), 'text': visible_text(p)})
            for name in sorted(z.namelist()):
                if name.startswith('word/') and name.endswith('.xml') and any(
                    name.startswith('word/'+prefix) for prefix in ['footnotes','endnotes','header','footer','comments']):
                    for i,p in enumerate(paragraphs(parse(z.read(name)))):
                        units.append({'location': f'{name} paragraph {i}', 'text': visible_text(p)})
            if root.findall('.//w:drawing',NS) or root.findall('.//w:pict',NS):
                warnings.append('Embedded visuals present; inspect them separately')
            if root.findall('.//w:del',NS) or root.findall('.//w:ins',NS):
                warnings.append('Tracked changes present; decide intended revision view before using extracted text')
        warnings.append('OMML extraction is a text index, not preservation of equation structure; inspect native equations')
    elif suffix == '.pptx':
        from pptx import Presentation
        def shape_text(shape):
            if shape.has_table:
                return '\n'.join('\t'.join(c.text for c in row.cells) for row in shape.table.rows)
            if hasattr(shape, 'shapes'): return '\n'.join(shape_text(s) for s in shape.shapes)
            return shape.text if shape.has_text_frame else ''
        prs = Presentation(path)
        for i,slide in enumerate(prs.slides,1):
            units.append({'location': f'slide {i}', 'text': '\n'.join(filter(None,(shape_text(s) for s in slide.shapes)))})
            if slide.has_notes_slide:
                frame = slide.notes_slide.notes_text_frame
                if frame is not None and frame.text.strip():
                    units.append({'location': f'slide {i} speaker notes', 'text': frame.text})
        warnings.append('Slide text order may differ from visual reading order; inspect charts, diagrams and equations')
    else:
        raise ValueError(f'Unsupported extension {suffix}; use visual/OCR or the appropriate format reader')
    return {'units': units, 'warnings': warnings}

def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__)
    ap.add_argument('inputs', nargs='+', type=Path)
    ap.add_argument('--output', required=True, type=Path)
    ap.add_argument('--reuse', action='store_true', help='Reuse matching hash/version from existing output')
    args = ap.parse_args(argv)
    if args.output.resolve() in [p.resolve() for p in args.inputs]:
        ap.error('Output must not overwrite an input')
    cached = {}
    if args.reuse and args.output.exists():
        old = json.loads(args.output.read_text())
        if old.get('extractor_version') == VERSION:
            cached = {s['path']:s for s in old.get('sources',[]) if 'error' not in s}
    report = {'extractor_version':VERSION, 'sources':[], 'errors':[]}
    for path in dict.fromkeys(p.resolve() for p in args.inputs):
        item = {'path':str(path)}
        try:
            item['sha256'] = sha256(path)
            prior = cached.get(str(path))
            if prior and prior.get('sha256') == item['sha256']:
                item.update(prior); item['cache_reused'] = True
            else:
                item.update(extract(path)); item['cache_reused'] = False
        except Exception as e:
            item['error'] = str(e)
            report['errors'].append({'path':str(path),'error':str(e)})
        report['sources'].append(item)
    write_json(args.output,report)
    print(json.dumps({'sources':len(report['sources']),'units':sum(len(s.get('units',[])) for s in report['sources']),
                      'errors':len(report['errors']),'output':str(args.output)},ensure_ascii=False))
    return 2 if report['errors'] else 0

if __name__ == '__main__': sys.exit(main())
